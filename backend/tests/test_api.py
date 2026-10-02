"""End-to-end API tests: upload → OCR (offline engine) → parse → review → export."""

from __future__ import annotations

import importlib
import io
import os
import time
from pathlib import Path

import cv2
import pytest
from docx import Document
from fastapi.testclient import TestClient

from tests import fixtures_gen

BEST = Path("/opt/tessdata_best")


@pytest.fixture(scope="module")
def client(tmp_path_factory):
    data_dir = tmp_path_factory.mktemp("data")
    env = {"DATA_DIR": str(data_dir), "ADMIN_PASSWORD": "", "ANTHROPIC_API_KEY": ""}
    if (BEST / "fas.traineddata").exists():
        env["TESSDATA_DIR"] = str(BEST)
    old = {k: os.environ.get(k) for k in env}
    os.environ.update(env)
    from app import config

    config.get_settings.cache_clear()
    main = importlib.reload(importlib.import_module("app.main"))
    with TestClient(main.app) as test_client:
        yield test_client
    for key, value in old.items():
        if value is None:
            os.environ.pop(key, None)
        else:
            os.environ[key] = value
    config.get_settings.cache_clear()


def _wait_ready(client: TestClient, project_id: str, timeout: float = 240) -> dict:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        project = client.get(f"/api/projects/{project_id}").json()
        if project["status"] in ("ready", "failed"):
            return project
        time.sleep(0.5)
    raise AssertionError("processing timed out")


def _upload(client: TestClient, files: list[tuple[str, bytes, str]], **form) -> dict:
    response = client.post(
        "/api/projects",
        files=[("booklet", f) for f in files],
        data={"track": "bar", "year": "1404", "engine": "offline", **form},
    )
    assert response.status_code == 200, response.text
    return response.json()


def _assert_four_questions(project: dict) -> None:
    assert project["status"] == "ready", project["error"]
    numbers = [q["number"] for q in project["questions"]]
    assert numbers == [1, 2, 3, 4]
    for question in project["questions"]:
        assert question["stem"]
        assert [o["key"] for o in question["options"]] == ["1", "2", "3", "4"]
        assert question["regions"], "each question must be locatable on the page"


def test_health_and_meta(client):
    health = client.get("/api/health").json()
    assert health["engines"]["offline"] is True
    assert health["push_configured"] is False
    meta = client.get("/api/meta").json()
    assert any(b["code"] == "BAR-1405" for b in meta["blueprints"])
    assert {"key": "civil", "name": "حقوق مدنی"} in meta["subjects"]


def test_typed_pdf_end_to_end(client):
    created = _upload(
        client,
        [("booklet.pdf", fixtures_gen.booklet_pdf(), "application/pdf")],
        blueprint="BAR-1405",
    )
    project = _wait_ready(client, created["id"])
    _assert_four_questions(project)
    assert project["questions"][0]["subject_key"] == "civil"

    pid = project["id"]
    page = client.get(f"/api/projects/{pid}/pages/booklet/0").json()
    assert page["source"] == "text_layer"
    image = client.get(f"/api/projects/{pid}/pages/booklet/0.jpg")
    assert image.status_code == 200 and image.headers["content-type"] == "image/jpeg"

    # edit + approve: manual key only when the key actually changes
    q1 = project["questions"][0]
    update = {
        "stem": q1["stem"],
        "options": q1["options"],
        "correct_key": "2",
        "explanation": "گزینه ۲ صحیح است.",
        "status": "approved",
    }
    saved = client.put(f"/api/projects/{pid}/questions/1", json=update).json()
    assert saved["status"] == "approved" and saved["key_source"] == "manual"
    assert not [i for i in saved["issues"] if i["code"] == "missing_key"]

    # add / delete
    assert client.post(f"/api/projects/{pid}/questions", json={"number": 5}).status_code == 200
    assert client.post(f"/api/projects/{pid}/questions", json={"number": 5}).status_code == 409
    assert client.delete(f"/api/projects/{pid}/questions/5").json() == {"ok": True}

    # exports
    payload = client.get(f"/api/projects/{pid}/export.json?only_approved=1").json()
    assert [q["source_number"] for q in payload["questions"]] == [1]
    assert payload["questions"][0]["correct_key"] == "2"
    docx = client.get(f"/api/projects/{pid}/export.docx")
    assert docx.status_code == 200
    text = "\n".join(p.text for p in Document(io.BytesIO(docx.content)).paragraphs)
    assert "سربرگ آزمون" in text and "سوال 4:" in text and "پاسخ صحیح: ب" in text

    # reparse rebuilds from stored OCR
    reparsed = client.post(f"/api/projects/{pid}/reparse", json={"blueprint": "auto"}).json()
    assert reparsed["blueprint"] == "auto" and len(reparsed["questions"]) == 4


def test_phone_photos_end_to_end(client):
    page = fixtures_gen.render(fixtures_gen.booklet_pdf(), dpi=200)
    photo = fixtures_gen.phone_scan(page)
    ok, jpeg = cv2.imencode(".jpg", photo, [cv2.IMWRITE_JPEG_QUALITY, 88])
    assert ok
    created = _upload(client, [("IMG_0001.jpg", jpeg.tobytes(), "image/jpeg")])
    assert created["documents"][0]["page_count"] == 1
    project = _wait_ready(client, created["id"])
    _assert_four_questions(project)
    result = client.get(f"/api/projects/{project['id']}/pages/booklet/0").json()
    assert result["source"] == "ocr" and result["preprocess"]


def test_multiple_images_become_pages(client):
    page = fixtures_gen.render(fixtures_gen.booklet_pdf(), dpi=120)
    _, png = cv2.imencode(".png", page)
    files = [("p1.png", png.tobytes(), "image/png"), ("p2.png", png.tobytes(), "image/png")]
    response = client.post(
        "/api/projects", files=[("booklet", f) for f in files], data={"engine": "offline"}
    )
    assert response.status_code == 200
    assert response.json()["documents"][0]["page_count"] == 2
    client.delete(f"/api/projects/{response.json()['id']}")


def test_rejects_invalid_upload(client):
    response = client.post(
        "/api/projects", files=[("booklet", ("x.pdf", b"not a pdf", "application/pdf"))]
    )
    assert response.status_code == 400
    assert "معتبر" in response.json()["detail"]


def test_unknown_project(client):
    assert client.get("/api/projects/doesnotexist").status_code == 404


def test_text_mode_end_to_end(client):
    created = _upload(
        client,
        [("notes.pdf", fixtures_gen.booklet_pdf(), "application/pdf")],
        doc_type="text",
    )
    assert created["mode"] == "text"
    project = _wait_ready(client, created["id"])
    assert project["status"] == "ready" and project["mode"] == "text"
    pid = project["id"]

    view = client.get(f"/api/projects/{pid}/pages/booklet/0/text").json()
    assert view["edited"] is False and view["approved"] is False
    assert "حقوق" in view["text"] or len(view["text"]) > 50

    edited = client.put(
        f"/api/projects/{pid}/pages/booklet/0/text",
        json={"text": "نکته‌ی اصلاح‌شده\nپاراگراف دوم", "approved": True},
    ).json()
    assert edited == {"text": "نکته‌ی اصلاح‌شده\nپاراگراف دوم", "edited": True, "approved": True}
    assert client.get(f"/api/projects/{pid}").json()["page_status"] == {"booklet:0": True}

    txt = client.get(f"/api/projects/{pid}/export.txt?only_approved=1")
    assert txt.text.strip() == "نکته‌ی اصلاح‌شده\n\nپاراگراف دوم".replace("\n\n", "\n")
    docx = client.get(f"/api/projects/{pid}/export-text.docx")
    paragraphs = [p.text for p in Document(io.BytesIO(docx.content)).paragraphs]
    assert "نکته‌ی اصلاح‌شده" in paragraphs and "پاراگراف دوم" in paragraphs

    # reset the edit, then switch to question mode (runs the question parser)
    reset = client.put(f"/api/projects/{pid}/pages/booklet/0/text", json={"text": None}).json()
    assert reset["edited"] is False
    switched = client.post(f"/api/projects/{pid}/mode", json={"mode": "questions"}).json()
    assert switched["mode"] == "questions" and len(switched["questions"]) == 4


def test_push_uploads_docx_to_site(client, monkeypatch):
    from app import site_import

    created = _upload(client, [("b.pdf", fixtures_gen.booklet_pdf(), "application/pdf")])
    pid = _wait_ready(client, created["id"])["id"]
    sent = {}

    def fake_upload(data, filename, settings):
        sent.update(size=len(data), filename=filename)
        return {"id": 7, "status": "queued"}

    monkeypatch.setattr(site_import, "upload_docx", fake_upload)
    assert client.post(f"/api/projects/{pid}/push", json={"only_approved": True}).status_code == 400
    response = client.post(f"/api/projects/{pid}/push", json={"only_approved": False}).json()
    assert response == {"ok": True, "questions": 4, "response": {"id": 7, "status": "queued"}}
    assert sent["filename"].endswith(".docx") and sent["size"] > 1000


def test_site_check_reports_missing_config(client):
    response = client.get("/api/site/check")
    assert response.status_code == 502
    assert "پیکربندی" in response.json()["detail"]


def test_batch_upload_auto_approve_queue_and_keys(client):
    pdf = fixtures_gen.booklet_pdf()
    response = client.post(
        "/api/projects/batch",
        files=[
            ("files", ("a.pdf", pdf, "application/pdf")),
            ("files", ("b.pdf", pdf, "application/pdf")),
        ],
        data={"doc_type": "questions", "engine": "offline", "auto_approve": "1"},
    )
    assert response.status_code == 200, response.text
    batch = response.json()
    assert len(batch["projects"]) == 2 and batch["errors"] == []
    assert {p["title"] for p in batch["projects"]} == {"a", "b"}
    assert all(p["batch_id"] == batch["batch_id"] and p["auto_approve"] for p in batch["projects"])
    a, b = (_wait_ready(client, p["id"]) for p in batch["projects"])
    assert client.get("/api/queue").json() == {"running": [], "queued": []}
    assert a["stats"]["pages"] == 1 and a["stats"]["engine"] == "text_layer"

    # the fixture has no answer keys → nothing is clean yet
    assert all(q["status"] == "pending" for q in a["questions"])
    queue = client.get(f"/api/review-queue?project_id={a['id']}").json()
    assert [i["number"] for i in queue] == [1, 2, 3, 4] and queue[0]["level"] == "error"
    assert "missing_key" in queue[0]["codes"]

    # quick key entry (Persian digits), then auto-approve the clean ones
    with_keys = client.put(f"/api/projects/{a['id']}/keys", json={"keys": "۲ ۴ ۱ ۳"}).json()
    assert [q["correct_key"] for q in with_keys["questions"]] == ["2", "4", "1", "3"]
    assert client.put(f"/api/projects/{a['id']}/keys", json={"keys": "25"}).status_code == 400
    result = client.post(f"/api/projects/{a['id']}/auto-approve").json()
    approved = [q for q in result["project"]["questions"] if q["status"] == "approved"]
    assert result["approved"] == len(approved)
    assert all(q["approved_by"] == "auto" for q in approved)

    # b is a copy of a: duplicates are detected across projects
    summaries = {s["id"]: s for s in client.get("/api/projects").json()}
    assert summaries[a["id"]]["auto_approved_count"] == result["approved"]
    if b["questions"] and any(q["duplicates"] for q in b["questions"]):
        assert summaries[b["id"]]["duplicate_count"] >= 1

    # editing an auto-approved question sends it back to review
    if approved:
        q = approved[0]
        edited = client.put(
            f"/api/projects/{a['id']}/questions/{q['number']}", json={"stem": q["stem"] + " "}
        ).json()
        assert edited["status"] == "pending" and edited["approved_by"] is None


def test_batch_reports_bad_files_without_failing(client):
    pdf = fixtures_gen.booklet_pdf()
    response = client.post(
        "/api/projects/batch",
        files=[
            ("files", ("ok.pdf", pdf, "application/pdf")),
            ("files", ("bad.pdf", b"nope", "application/pdf")),
        ],
        data={"engine": "offline"},
    )
    body = response.json()
    assert len(body["projects"]) == 1 and body["errors"][0]["filename"] == "bad.pdf"
    _wait_ready(client, body["projects"][0]["id"])


def test_batch_titles_zip_export_and_dismissed_duplicates(client):
    import zipfile

    pdf = fixtures_gen.booklet_pdf()
    body = client.post(
        "/api/projects/batch",
        files=[
            ("files", ("x.pdf", pdf, "application/pdf")),
            ("files", ("y.pdf", pdf, "application/pdf")),
        ],
        data={"engine": "offline", "doc_type": "questions", "titles": ["دفترچه الف", ""]},
    ).json()
    assert [p["title"] for p in body["projects"]] == ["دفترچه الف", "y"]
    first, second = (_wait_ready(client, p["id"]) for p in body["projects"])
    summary = next(s for s in client.get("/api/projects").json() if s["id"] == first["id"])
    assert summary["stats"]["pages"] == 1

    # duplicates: the second copy sees the first; «تکراری نیست» clears and persists
    dup = next((q for q in second["questions"] if q["duplicates"]), None)
    if dup is not None:
        cleared = client.put(
            f"/api/projects/{second['id']}/questions/{dup['number']}", json={"duplicates": []}
        ).json()
        assert cleared["duplicates"] == [] and "duplicate" not in {
            i["code"] for i in cleared["issues"]
        }

    # zip: approved only → nothing yet; all → one docx per project
    ids = f"{first['id']},{second['id']}"
    assert client.get(f"/api/export.zip?project_ids={ids}").status_code == 400
    response = client.get(f"/api/export.zip?project_ids={ids}&only_approved=false")
    assert response.headers["content-type"] == "application/zip"
    names = zipfile.ZipFile(io.BytesIO(response.content)).namelist()
    assert sorted(names) == ["dafterche.docx"] or sorted(names) == sorted(
        ["دفترچه الف.docx", "y.docx"]
    )
